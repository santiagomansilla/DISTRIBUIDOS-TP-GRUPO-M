import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_manual():
    doc = Document()

    # Configurar márgenes
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Portada / Encabezado principal
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("Universidad Nacional de Lanús\n")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(18)
    title_run.bold = True
    title_run.font.color.rgb = RGBColor(30, 60, 114)

    sub_run = title_p.add_run("Desarrollo de Software en Sistemas Distribuidos\n")
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(14)
    sub_run.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_paragraph()

    h1_p = doc.add_paragraph()
    h1_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1_run = h1_p.add_run("MANUAL DE INSTALACIÓN, EJECUCIÓN Y USO\nSISTEMA RENTAR (HITOS 1 Y 2)")
    h1_run.font.name = "Arial"
    h1_run.font.size = Pt(16)
    h1_run.bold = True
    h1_run.font.color.rgb = RGBColor(42, 82, 152)

    doc.add_paragraph("─" * 45).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 1. Introducción
    doc.add_heading("1. Descripción del Proyecto", level=1)
    p = doc.add_paragraph(
        "El proyecto Rentar es una solución desarrollada para la administración integral de una empresa de alquiler de vehículos. "
        "Fue concebido con una arquitectura distribuida que implementa múltiples mecanismos de comunicación sincrónica y asincrónica "
        "a lo largo de tres hitos evolutivos:\n"
        "• Hito 1: Servicios web mediante REST (con documentación Swagger/OpenAPI) y GraphQL para consultas flexibles.\n"
        "• Hito 2: Comunicación RPC de alto rendimiento utilizando gRPC y Protocol Buffers (.proto).\n"
        "• Hito 3: Integración de eventos y mensajería distribuida (Kafka / RabbitMQ)."
    )

    # 2. Requisitos Previos
    doc.add_heading("2. Requisitos Previos", level=1)
    doc.add_paragraph(
        "Para ejecutar el proyecto en su máquina local se requiere contar con:\n"
        "1. Java Development Kit (JDK): Versión 21 instalada y configurada en las variables de entorno.\n"
        "2. Acceso a Internet: Durante la primera compilación para la descarga automática de librerías y dependencias vía Maven Wrapper.\n"
        "3. Navegador Web moderno: Chrome, Edge, Firefox o Safari para interactuar con la interfaz web, Swagger y GraphiQL."
    )

    # 3. Cómo Levantar el Proyecto
    doc.add_heading("3. Pasos para Compilar y Levantar el Proyecto", level=1)
    doc.add_paragraph(
        "El proyecto incluye el empaquetador Maven Wrapper (mvnw.cmd), por lo que no es necesario instalar Maven de forma manual."
    )
    
    doc.add_paragraph("Paso 1: Abrir la terminal", style='List Bullet')
    p_code1 = doc.add_paragraph()
    r = p_code1.add_run("cd \"C:\\Users\\aggus\\OneDrive\\Desktop\\Sistemas Distribuidos\\tp-distribuidos\"")
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)

    doc.add_paragraph("Paso 2: Liberar el puerto 8080 si existe una instancia previa", style='List Bullet')
    doc.add_paragraph("Asegúrese de cerrar cualquier terminal previa de Java o proceso que esté escuchando en el puerto 8080.")

    doc.add_paragraph("Paso 3: Ejecutar la aplicación", style='List Bullet')
    p_code2 = doc.add_paragraph()
    r2 = p_code2.add_run(".\\mvnw.cmd spring-boot:run")
    r2.font.name = "Consolas"
    r2.font.size = Pt(9.5)
    r2.bold = True

    doc.add_paragraph(
        "Cuando la aplicación haya iniciado con éxito, observará el mensaje de Spring Boot confirmando el inicio del servidor web en el puerto 8080 y el servidor gRPC en el puerto 9090."
    )

    # 4. Uso de la Aplicación y Puntos de Acceso
    doc.add_heading("4. Puntos de Entrada y Uso de la Aplicación", level=1)

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Componente"
    hdr_cells[1].text = "URL / Puerto"
    hdr_cells[2].text = "Descripción"

    items = [
        ("Interfaz Web Principal", "http://localhost:8080", "Pantalla de operaciones para clientes y administradores (búsqueda, reservas, historial, ABM)."),
        ("Swagger UI (REST)", "http://localhost:8080/swagger-ui/index.html", "Documentación interactiva de todos los endpoints REST desarrollados."),
        ("GraphiQL Console", "http://localhost:8080/graphiql", "Consola interactiva para probar consultas GraphQL con autocompletado y validación."),
        ("Base de Datos H2", "http://localhost:8080/h2-console", "Consola web de H2 (JDBC URL: jdbc:h2:mem:tpdb, User: sa, Password en blanco)."),
        ("Servidor gRPC", "localhost:9090", "Servicio RPC de alto rendimiento implementado con Protocol Buffers.")
    ]

    for comp, url, desc in items:
        row_cells = table.add_row().cells
        row_cells[0].text = comp
        row_cells[1].text = url
        row_cells[2].text = desc

    doc.add_paragraph()

    # 5. Funcionalidades del Sistema
    doc.add_heading("5. Detalle de Operaciones y Flujo de Trabajo", level=1)
    doc.add_paragraph(
        "A continuación se describe el funcionamiento de cada requerimiento del Trabajo Práctico:\n\n"
        "• Requerimiento 1 - Gestión de Vehículos (REST): Permite el alta, modificación, consulta y baja lógica de vehículos. Al registrar un nuevo vehículo, éste ingresa automáticamente en estado DISPONIBLE. La patente es única y no modificable.\n\n"
        "• Requerimiento 2 - Consulta de Disponibilidad (GraphQL): Permite buscar vehículos filtrando por fechas obligatorias de inicio y fin. El motor verifica en la base de datos que el vehículo no posea reservas solapadas en ese período.\n\n"
        "• Requerimiento 3 - Gestión de Clientes (REST): ABM completo de clientes asegurando unicidad en documento y correo electrónico, con baja lógica para inhabilitar nuevos alquileres.\n\n"
        "• Requerimiento 4 - Alta de Reservas (REST): Realiza la reserva validando que cliente y vehículo estén activos, que las fechas sean futuras y congruentes, y calcula el importe total en función de los días y la tarifa diaria.\n\n"
        "• Requerimiento 5 - Consulta de Reservas (GraphQL): Soporta filtros opcionales de cliente, vehículo, fechas y estado, diferenciando consultas de cliente vs administrador.\n\n"
        "• Requerimiento 6 - Cancelación de Reserva (REST): Anula una reserva modificando su estado a CANCELADA, liberando inmediatamente las fechas del vehículo.\n\n"
        "• Requerimiento 7 - Historial de Alquileres (GraphQL): Expone el histórico de alquileres finalizados y cancelados para cada cliente."
    )

    # 6. Datos de Prueba Precargados
    doc.add_heading("6. Datos de Prueba Precargados (DataLoader)", level=1)
    doc.add_paragraph(
        "Al iniciar el sistema por primera vez en memoria, se generan automáticamente los siguientes datos:\n"
        "• Clientes:\n"
        "   - ID 1: Agustín González (Doc: 38450123, Email: agustin@unla.edu.ar)\n"
        "   - ID 2: María Fernández (Doc: 40123987, Email: maria.fernandez@gmail.com)\n"
        "• Vehículos:\n"
        "   - ID 1: Toyota Corolla (SEDAN, Patente: AE123CD, $45.000/día)\n"
        "   - ID 2: Ford Ranger (PICKUP, Patente: AF456GH, $75.000/día)\n"
        "   - ID 3: Volkswagen Taos (SUV, Patente: AD789JK, $60.000/día)\n"
        "   - ID 4: Peugeot 208 (HATCHBACK, Patente: AC321MN, $38.000/día)"
    )

    # Guardar documento
    output_path = r"C:\Users\aggus\OneDrive\Desktop\Sistemas Distribuidos\tp-distribuidos\Manual_Instalacion_y_Uso_Rentar.docx"
    doc.save(output_path)
    print("Documento creado exitosamente en:", output_path)

if __name__ == "__main__":
    create_manual()
